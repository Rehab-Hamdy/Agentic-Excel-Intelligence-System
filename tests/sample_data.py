"""
Generate a realistic sample sales.xlsx dataset for testing.
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

# Ensure the project root is on the path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


def generate_sales_data(num_rows: int = 200, seed: int = 42) -> pd.DataFrame:
    """Generate a realistic sales dataset."""
    rng = np.random.default_rng(seed)

    products = ["Laptop", "Phone", "Tablet", "Monitor", "Keyboard",
                "Mouse", "Headphones", "Webcam", "Charger", "Cable"]
    regions = ["North", "South", "East", "West"]
    categories = {
        "Laptop": "Electronics", "Phone": "Electronics", "Tablet": "Electronics",
        "Monitor": "Electronics", "Keyboard": "Accessories", "Mouse": "Accessories",
        "Headphones": "Audio", "Webcam": "Peripherals", "Charger": "Accessories",
        "Cable": "Accessories",
    }

    dates = pd.date_range(start="2024-01-01", periods=365, freq="D")
    chosen_dates = rng.choice(dates, size=num_rows, replace=True)
    chosen_products = rng.choice(products, size=num_rows)
    chosen_regions = rng.choice(regions, size=num_rows)

    # Price ranges by product
    price_map = {
        "Laptop": (800, 2000), "Phone": (400, 1200), "Tablet": (300, 900),
        "Monitor": (200, 800), "Keyboard": (30, 150), "Mouse": (15, 80),
        "Headphones": (50, 300), "Webcam": (40, 200), "Charger": (15, 60),
        "Cable": (5, 30),
    }

    unit_prices = [round(rng.uniform(*price_map[p]), 2) for p in chosen_products]
    quantities = rng.integers(1, 20, size=num_rows)
    revenues = [round(up * q, 2) for up, q in zip(unit_prices, quantities)]
    costs = [round(r * rng.uniform(0.4, 0.75), 2) for r in revenues]
    profits = [round(r - c, 2) for r, c in zip(revenues, costs)]

    df = pd.DataFrame({
        "Date": sorted(chosen_dates),
        "Product": chosen_products,
        "Category": [categories[p] for p in chosen_products],
        "Region": chosen_regions,
        "Unit_Price": unit_prices,
        "Quantity": quantities,
        "Revenue": revenues,
        "Cost": costs,
        "Profit": profits,
    })

    # Inject a few missing values for realism
    missing_indices = rng.choice(len(df), size=5, replace=False)
    for idx in missing_indices:
        df.loc[idx, "Unit_Price"] = np.nan

    return df


def save_as_excel(df: pd.DataFrame, file_path: Path):
    """Save DataFrame to a nicely formatted Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

    # Auto-width
    for col_cells in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 30)

    wb.save(str(file_path))
    wb.close()
    print(f"[OK] Created: {file_path} ({len(df)} rows)")


if __name__ == "__main__":
    output_dir = Path(__file__).resolve().parent.parent / "data" / "input"
    output_dir.mkdir(parents=True, exist_ok=True)

    df = generate_sales_data()
    save_as_excel(df, output_dir / "sales.xlsx")

    print("\nSample columns:", list(df.columns))
    print("Shape:", df.shape)
    print("\nFirst 5 rows:")
    print(df.head().to_string(index=False))
